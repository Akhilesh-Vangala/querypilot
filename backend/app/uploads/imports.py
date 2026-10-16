"""Bounded tabular imports; never execute spreadsheet formulas or arbitrary files."""
import base64
import binascii
import csv
from datetime import date, datetime, time
from decimal import Decimal
import hashlib
import io
import json
from pathlib import PurePosixPath
import zipfile

from pydantic import BaseModel, ConfigDict, Field
from .ingestion import CSV_LIMIT, MAX_ROWS, UploadProblem, profile

FORMATS = {"csv", "tsv", "json", "jsonl", "xlsx"}

class FileImport(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    filename: str = Field(min_length=1, max_length=160)
    content_base64: str = Field(min_length=1, max_length=533336)
    sheet: str | None = Field(default=None, max_length=120)


def refuse(message, code="invalid_import", **kwargs):
    raise UploadProblem(message, code=code, hint="Choose a supported tabular export, check the indicated issue and retry.", **kwargs)


def scalar(value):
    if value is None: return ""
    if type(value) is bool: return "true" if value else "false"
    if isinstance(value, (datetime, date, time)): return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        number=Decimal(str(value))
        if not number.is_finite() or abs(number.adjusted())>100 or len(number.as_tuple().digits)>100:
            refuse("A numeric value exceeds supported import precision or magnitude")
        return format(number,"f")
    if isinstance(value,str): return value
    refuse("Nested objects and arrays are not supported as table cells")


def table_csv(headers, rows, max_rows):
    if not headers or any(not isinstance(h,str) for h in headers): refuse("Use one row of named text headers")
    if len(rows)>max_rows: refuse(f"This app supports at most {max_rows:,} data records per file")
    text=io.StringIO(newline="");writer=csv.writer(text,lineterminator="\n")
    writer.writerow(headers)
    for number,row in enumerate(rows,2):
        if len(row)!=len(headers): refuse("Records must use a consistent column count",issues=[{"row":number,"message":"The record does not match the header width."}])
        writer.writerow([scalar(v) for v in row])
        if text.tell()>CSV_LIMIT: refuse("The normalized table exceeds 400 KB")
    value=text.getvalue()
    if len(value.encode())>CSV_LIMIT: refuse("The normalized table exceeds 400 KB")
    return value


def records_json(text, line_mode, max_rows):
    def pairs(items):
        result={}
        for key,value in items:
            if key in result: refuse("Duplicate JSON object keys are ambiguous",code="duplicate_headers")
            result[key]=value
        return result
    def constant(_):refuse("Non-finite JSON numbers are unsupported")
    def decode(value):return json.loads(value,parse_float=Decimal,parse_int=Decimal,parse_constant=constant,object_pairs_hook=pairs)
    try:
        if line_mode:
            lines=text.splitlines()
            if any(not line.strip() for line in lines):refuse("JSONL must have one object on each non-empty line; remove blank records")
            if len(lines)>max_rows:refuse(f"This app supports at most {max_rows:,} data records per file")
            records=[decode(line) for line in lines]
        else: records=decode(text)
    except (json.JSONDecodeError,RecursionError,ValueError) as error:
        if isinstance(error,UploadProblem):raise
        refuse("Invalid JSON; use an array of flat objects or one object per JSONL line")
    if not isinstance(records,list) or not records or len(records)>max_rows or not all(isinstance(r,dict) for r in records):
        refuse(f"Use 1–{max_rows:,} flat record objects")
    headers=list(records[0]);rows=[]
    for number,record in enumerate(records,2):
        if set(record)!=set(headers):
            refuse("Every JSON record must have the same fields; use null for a missing value",issues=[{"row":number,"message":"Missing or additional fields differ from the first record."}])
        rows.append([record[h] for h in headers])
    return table_csv(headers,rows,max_rows)


def excel(raw, sheet, max_rows):
    from defusedxml import ElementTree as SafeXML
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_to_tuple
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries=archive.infolist();names=[i.filename for i in entries]
            if len(entries)>200 or len(set(names))!=len(names):refuse("Workbook structure exceeds supported bounds")
            if sum(i.file_size for i in entries)>8_000_000:refuse("Workbook expands beyond the supported 8 MB limit")
            for info in entries:
                path=PurePosixPath(info.filename)
                if info.flag_bits&1 or path.is_absolute() or '..' in path.parts or info.file_size>4_000_000:
                    refuse("Encrypted or oversized workbook content is unsupported")
                if info.file_size>max(1,info.compress_size)*150:refuse("Workbook compression ratio exceeds the supported bound")
                if any(t in info.filename.lower() for t in ('vbaproject','externallinks/','embeddings/','activex/')):
                    refuse("Macros, embedded objects and external workbook links are unsupported")
                if info.filename.endswith(('.xml','.rels')):
                    content=archive.read(info)
                    if b'macroenabled' in content.lower():refuse("Macro-enabled workbooks are unsupported")
                    document=SafeXML.fromstring(content,forbid_dtd=True)
                    if info.filename.startswith('xl/worksheets/') and info.filename.endswith('.xml'):
                        for node in document.iter():
                            tag=node.tag.rsplit('}',1)[-1]
                            if tag=='mergeCell':refuse("Merged cells are ambiguous; unmerge them before importing")
                            if tag=='f':refuse("Spreadsheet formulas are unsupported; export a values-only copy")
                            if tag=='c':
                                row,col=coordinate_to_tuple(node.get('r','A1'))
                                if row>MAX_ROWS+1 or col>60:refuse("Workbook cells exceed 5,000 data records or 60 columns")
            if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:refuse("The file is not a valid .xlsx workbook")
        workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        try:
            sheets=[s.title for s in workbook.worksheets if s.sheet_state=='visible']
            if not 1<=len(sheets)<=20:refuse("Choose a workbook with 1–20 visible worksheets")
            if sheet is None and len(sheets)>1:return {"needs_sheet":True,"sheets":sheets}
            selected=sheet or sheets[0]
            if selected not in sheets:refuse("Choose one of the workbook's visible worksheets")
            ws=workbook[selected]
            # Do not trust dimension metadata to truncate the actual cell stream.
            ws.reset_dimensions()
            rows=[];width=0
            for number,cells in enumerate(ws.iter_rows(),1):
                if number>max_rows+1:refuse(f"Worksheet exceeds {max_rows:,} data records")
                values=[c.value for c in cells]
                if len(values)>60:refuse("Worksheet exceeds 60 columns")
                if number==1:
                    while values and values[-1] is None:values.pop()
                    width=len(values)
                elif len(values)>width and any(v is not None for v in values[width:]):refuse("A worksheet record extends past the header width")
                values=(values+[None]*width)[:width]
                rows.append(values)
            if not rows:refuse("The selected worksheet has no table")
            return {"csv":table_csv(rows[0],rows[1:],max_rows),"sheet":selected}
        finally:workbook.close()
    except UploadProblem:raise
    except Exception:
        # Parser exceptions may contain cell values/paths: publish fixed guidance.
        refuse("The workbook could not be read; use an unencrypted values-only .xlsx export")


def import_file(body, *, max_rows=MAX_ROWS):
    extension=body.filename.rsplit('.',1)[-1].lower()
    if extension not in FORMATS:refuse("Supported formats: CSV, TSV, Excel .xlsx, JSON and JSONL",code="unsupported_format")
    try:raw=base64.b64decode(body.content_base64,validate=True)
    except (binascii.Error,ValueError):refuse("Invalid file transfer; choose the file again")
    if not raw or len(raw)>CSV_LIMIT:refuse("Choose a non-empty file no larger than 400 KB")
    if extension!='xlsx' and body.sheet is not None:refuse("Worksheet selection applies only to Excel workbooks")
    origin={"format":extension,"original_sha256":hashlib.sha256(raw).hexdigest(),"original_bytes":len(raw),"normalization":"CSV retained verbatim; other formats converted to a named rectangular table. Review all mappings before execution."}
    if extension=='xlsx':
        result=excel(raw,body.sheet,max_rows)
        if result.get('needs_sheet'):return {**result,"source":origin}
        source=result['csv'];origin['sheet']=result['sheet']
    else:
        try:text=raw.decode('utf-8' if extension=='csv' else 'utf-8-sig')
        except UnicodeError:refuse("Use UTF-8 text or an Excel .xlsx workbook")
        if extension=='csv':source=text
        elif extension=='tsv':
            try:
                records=list(csv.reader(io.StringIO(text,newline=''),delimiter='\t',strict=True))
                if not records:refuse("TSV has no header")
                source=table_csv(records[0],records[1:],max_rows)
            except csv.Error:refuse("Invalid TSV quoting or field separators")
        else:source=records_json(text,extension=='jsonl',max_rows)
    inspected=profile(source)
    if inspected['rows']>max_rows:refuse(f"This app supports at most {max_rows:,} data records per file")
    origin['normalized_sha256']=inspected['sha256']
    return {"needs_sheet":False,"csv":source,"profile":inspected,"source":origin}
