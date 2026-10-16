"""Actual format conversion, adversarial files and preservation of business results."""
import base64
import copy
import csv
import hashlib
import io
import json
from pathlib import Path
import unittest
import zipfile
from openpyxl import Workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient
from app.uploads.api import attach
from app.uploads.imports import FileImport, import_file
from app.uploads.ingestion import Upload, UploadProblem, parse
from app.uploads.service import run

class MultiFormatTests(unittest.TestCase):
    def body(self,content,name="data.json",sheet=None):
        raw=content.encode() if isinstance(content,str) else content
        return FileImport(filename=name,content_base64=base64.b64encode(raw).decode(),sheet=sheet)

    def workbook(self,second=False,formula=False):
        book=Workbook();sheet=book.active;sheet.title="Records"
        sheet.append(["id","amount","confirmed"]);sheet.append(["00017","=2+2" if formula else 0.10,True]);sheet.append(["00018",0.20,False])
        if second:
            other=book.create_sheet("Other table");other.append(["id","amount"]);other.append(["different",99])
        stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()

    def test_utf8_csv_bom_and_crlf_remain_byte_identical_for_saved_source_restore(self):
        source="\ufeffid,amount\r\nA,1.00\r\n"
        result=import_file(self.body(source,"source.csv"))
        self.assertEqual(result["csv"],source)
        self.assertEqual(result["profile"]["sha256"],hashlib.sha256(source.encode()).hexdigest())

    def test_tsv_quoted_cells_and_unicode_are_rectangular_without_numeric_guessing(self):
        result=import_file(self.body('id\tamount\tnote\n00017\t0.10\t"é,quoted\tcell"\n',"file.tsv"))
        headers,rows=parse(result["csv"])
        self.assertEqual(headers,["id","amount","note"])
        self.assertEqual(rows,[["00017","0.10","é,quoted\tcell"]])

    def test_json_and_jsonl_preserve_decimal_lexemes_null_bool_and_field_order(self):
        for source,name in [('[{"id":"00017","value":0.100000,"flag":true,"empty":null},{"empty":null,"flag":false,"value":2e-6,"id":"00018"}]',"rows.json"),('{"id":"00017","value":0.100000,"flag":true,"empty":null}\n{"empty":null,"flag":false,"value":2e-6,"id":"00018"}\n',"rows.jsonl")]:
            result=import_file(self.body(source,name));headers,rows=parse(result["csv"])
            self.assertEqual(headers,["id","value","flag","empty"])
            self.assertEqual(rows,[["00017","0.100000","true",""],["00018","0.000002","false",""]])
            self.assertEqual(result["source"]["format"],name.rsplit('.',1)[-1])

    def test_ambiguous_nested_nonfinite_and_inconsistent_json_are_refused(self):
        cases=['[{"id":"a","id":"b"}]','[{"id":1,"ID":2}]','[{"id":{"nested":1}}]','[{"id":[1]}]','[{"id":1},{"id":2,"extra":3}]','[{"id":NaN}]','[{"id":1e999999}]','{"records":[]}','[]']
        for source in cases:
            with self.subTest(source=source),self.assertRaises(ValueError):import_file(self.body(source))

    def test_excel_values_and_explicit_visible_worksheet_selection(self):
        raw=self.workbook(second=True)
        choice=import_file(self.body(raw,"book.xlsx"))
        self.assertTrue(choice["needs_sheet"]);self.assertEqual(choice["sheets"],["Records","Other table"]);self.assertNotIn("csv",choice)
        result=import_file(self.body(raw,"book.xlsx",sheet="Records"))
        self.assertEqual(parse(result["csv"])[1],[["00017","0.1","true"],["00018","0.2","false"]])
        self.assertEqual(result["source"]["sheet"],"Records")
        with self.assertRaises(UploadProblem):import_file(self.body(raw,"book.xlsx",sheet="Missing"))

    def test_excel_formulas_merged_headers_macros_entities_and_zip_bombs_are_refused(self):
        with self.assertRaisesRegex(ValueError,"formulas"):import_file(self.body(self.workbook(formula=True),"formula.xlsx"))
        raw=self.workbook()
        def edited(transform):
            output=io.BytesIO()
            with zipfile.ZipFile(io.BytesIO(raw)) as source,zipfile.ZipFile(output,"w",zipfile.ZIP_DEFLATED) as target:
                for info in source.infolist():target.writestr(info.filename,transform(info.filename,source.read(info)))
            return output.getvalue()
        for bad in [edited(lambda name,data: data.replace(b'<worksheet ',b'<!DOCTYPE worksheet [<!ENTITY x "boom">]><worksheet ') if name=='xl/worksheets/sheet1.xml' else data),edited(lambda name,data:data.replace(b'</worksheet>',b'<mergeCells><mergeCell ref="A1:B1"/></mergeCells></worksheet>') if name=='xl/worksheets/sheet1.xml' else data)]:
            with self.assertRaises(UploadProblem):import_file(self.body(bad,"bad.xlsx"))
        output=io.BytesIO()
        with zipfile.ZipFile(output,"w",zipfile.ZIP_DEFLATED) as archive:archive.writestr('bomb.xml',b'x'*1000000)
        with self.assertRaises(UploadProblem):import_file(self.body(output.getvalue(),"bomb.xlsx"))

    def test_lied_workbook_dimensions_cannot_truncate_source_rows(self):
        raw=self.workbook();output=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as source,zipfile.ZipFile(output,"w",zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():target.writestr(info.filename,source.read(info).replace(b'ref="A1:C3"',b'ref="A1:A1"'))
        self.assertEqual(import_file(self.body(output.getvalue(),"table.xlsx"))["profile"]["rows"],2)

    def test_format_body_row_and_normalized_size_limits_are_enforced(self):
        for content,name in [(b'x','file.pdf'),(b'not a workbook','file.xlsx'),(b'x'*400001,'large.csv'),(b'','empty.csv'),(b'\xff','invalid.tsv')]:
            with self.assertRaises(ValueError):import_file(self.body(content,name))
        with self.assertRaises(ValueError):import_file(self.body('[{"a":1},{"a":2}]'),max_rows=1)
        with self.assertRaises(ValueError):import_file(FileImport(filename='x.csv',content_base64='!!!'))

    def test_actual_import_api_profile_structured_errors_and_request_boundary(self):
        app=FastAPI();attach(app,"QueryPilot",Path(__file__).resolve().parents[3])
        with TestClient(app) as client:
            result=client.post('/api/uploads/import',json=self.body('[{"id":"a","value":1.25}]').model_dump())
            self.assertEqual(result.status_code,200,result.text);self.assertEqual(result.headers['cache-control'],'no-store')
            self.assertEqual(result.json()['profile']['rows'],1)
            self.assertEqual(client.post('/api/uploads/import',json=self.body('[{"id":1,"id":2}]').model_dump()).status_code,422)
            self.assertEqual(client.post('/api/uploads/import',content=iter([b'x'*600000,b'x'*600000]),headers={'content-type':'application/json'}).status_code,413)
            self.assertEqual(client.post('/api/uploads/import',content='x',headers={'content-type':'text/csv'}).status_code,415)

    def test_all_supported_formats_feed_same_existing_domain_engine(self):
        source="region,amount,status\nEast,0.10,done\nEast,0.20,done\nWest,9.90,pending\n"
        mapping={"group":"c0","measure":"c1","filter":"c2"};options={"operation":"sum","filter_value":"done"}
        headers,rows=parse(source);objects=[dict(zip(headers,row)) for row in rows]
        tsv=io.StringIO();writer=csv.writer(tsv,delimiter='\t',lineterminator='\n');writer.writerow(headers);writer.writerows(rows)
        book=Workbook();sheet=book.active;sheet.append(headers)
        for row in rows:sheet.append(row)
        stream=io.BytesIO();book.save(stream);book.close()
        baseline=run(Upload(csv=source,mapping=mapping,options=options))
        for content,name in [(source,'data.csv'),(tsv.getvalue(),'data.tsv'),(json.dumps(objects),'data.json'),('\n'.join(json.dumps(r) for r in objects),'data.jsonl'),(stream.getvalue(),'data.xlsx')]:
            with self.subTest(format=name):
                imported=import_file(self.body(content,name));result=run(Upload(csv=imported['csv'],mapping=mapping,options=options))
                self.assertEqual(result['summary'],baseline['summary'])
                self.assertTrue(all(result['checks'].values()))
